# -*- coding: utf-8 -*-
from docx import Document
import sys
import io

doc_path = r'C:\Users\23145\Desktop\作品集图片\作品集图片.docx'
doc = Document(doc_path)

output_path = r'd:\PycharmProjects（D盘）\intelligent-audit-system\docx_output.txt'
with open(output_path, 'w', encoding='utf-8') as f:
    f.write("=== 段落内容 ===\n\n")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            f.write(f"[段落{i}] {text}\n")
    
    f.write("\n\n=== 表格内容 ===\n\n")
    for i, table in enumerate(doc.tables):
        f.write(f"\n--- 表格 {i+1} ---\n")
        for row_idx, row in enumerate(table.rows):
            row_text = [cell.text.strip() for cell in row.cells]
            f.write(f"行{row_idx}: {' | '.join(row_text)}\n")
    
    # 打印文档结构信息
    f.write("\n\n=== 文档结构 ===\n")
    f.write(f"段落数量: {len(doc.paragraphs)}\n")
    f.write(f"表格数量: {len(doc.tables)}\n")
    
    # 提取所有run的文本（可能包含不同格式的文本）
    f.write("\n\n=== 所有文本内容 ===\n")
    for para in doc.paragraphs:
        for run in para.runs:
            if run.text.strip():
                f.write(f"{run.text}\n")

print(f"内容已保存到: {output_path}")
