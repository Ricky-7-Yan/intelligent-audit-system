"""Build a recruiter-friendly Markdown/TXT/Word/PDF project archive."""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOC_FILES = sorted((PROJECT_ROOT / "docs").glob("[0-9][0-9]-*.md"))


def plain_text(markdown: str) -> str:
    text = re.sub(r"```[\w-]*\n", "", markdown)
    text = text.replace("```", "")
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text


def add_markdown_to_docx(document: Document, markdown: str) -> None:
    in_code = False
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if not line:
            document.add_paragraph()
            continue
        if in_code:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
            continue
        if line.startswith("- "):
            document.add_paragraph(plain_text(line[2:]), style="List Bullet")
            continue
        if re.match(r"^\d+\.\s+", line):
            document.add_paragraph(plain_text(re.sub(r"^\d+\.\s+", "", line)), style="List Number")
            continue
        document.add_paragraph(plain_text(line))


def build_docx(contents: list[tuple[str, str]], target: Path) -> None:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    title = document.add_heading("审脉 AuditPilot 项目完整说明", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("企业级智能审计 Agent · v4.0")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()
    for index, (_, markdown) in enumerate(contents):
        add_markdown_to_docx(document, markdown)
        if index < len(contents) - 1:
            document.add_page_break()
    document.save(target)


def build_pdf(contents: list[tuple[str, str]], target: Path) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ChineseBody",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=9.2,
        leading=14,
        spaceAfter=5,
        alignment=TA_LEFT,
    )
    headings = {
        1: ParagraphStyle("H1CN", parent=body, fontSize=18, leading=24, spaceBefore=10, spaceAfter=10),
        2: ParagraphStyle("H2CN", parent=body, fontSize=14, leading=19, spaceBefore=8, spaceAfter=7),
        3: ParagraphStyle("H3CN", parent=body, fontSize=11.5, leading=16, spaceBefore=6, spaceAfter=5),
    }
    title_style = ParagraphStyle(
        "TitleCN",
        parent=body,
        fontSize=24,
        leading=32,
        alignment=TA_CENTER,
        spaceAfter=14,
    )
    story = [
        Spacer(1, 42 * mm),
        Paragraph("审脉 AuditPilot 项目完整说明", title_style),
        Paragraph("企业级智能审计 Agent · v4.0", ParagraphStyle("SubCN", parent=body, fontSize=12, alignment=TA_CENTER)),
        PageBreak(),
    ]
    for file_index, (_, markdown) in enumerate(contents):
        in_code = False
        for raw_line in markdown.splitlines():
            line = raw_line.strip()
            if line.startswith("```"):
                in_code = not in_code
                continue
            if not line:
                story.append(Spacer(1, 3 * mm))
                continue
            heading = re.match(r"^(#{1,3})\s+(.+)$", line)
            if heading:
                story.append(Paragraph(plain_text(heading.group(2)), headings[len(heading.group(1))]))
            else:
                prefix = "• " if line.startswith("- ") else ""
                cleaned = plain_text(line[2:] if prefix else line)
                escaped = cleaned.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(prefix + escaped, body))
        if file_index < len(contents) - 1:
            story.append(PageBreak())
    pdf = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="审脉 AuditPilot 项目完整说明",
        author="AuditPilot",
    )
    pdf.build(story)


def build(destination: Path) -> None:
    if destination.exists():
        shutil.rmtree(destination)
    markdown_dir = destination / "Markdown"
    text_dir = destination / "clean_txt"
    word_dir = destination / "Word"
    pdf_dir = destination / "PDF"
    screenshot_dir = destination / "screenshots"
    jd_dir = destination / "JD来源资料"
    interview_dir = destination / "面经资料"
    for directory in (markdown_dir, text_dir, word_dir, pdf_dir, screenshot_dir):
        directory.mkdir(parents=True, exist_ok=True)
    jd_dir.mkdir(parents=True, exist_ok=True)
    interview_dir.mkdir(parents=True, exist_ok=True)

    contents = []
    for source in DOC_FILES:
        markdown = source.read_text(encoding="utf-8")
        contents.append((source.name, markdown))
        shutil.copy2(source, markdown_dir / source.name)
        (text_dir / source.with_suffix(".txt").name).write_text(plain_text(markdown), encoding="utf-8")

    readme = PROJECT_ROOT / "README.md"
    shutil.copy2(readme, destination / "README_归档说明.md")
    source_jd_dir = PROJECT_ROOT / "docs" / "jd_research"
    if source_jd_dir.exists():
        for source in source_jd_dir.glob("*.md"):
            shutil.copy2(source, jd_dir / source.name)
    source_interview_dir = PROJECT_ROOT / "docs" / "interview_experience"
    if source_interview_dir.exists():
        for source in source_interview_dir.glob("*.md"):
            shutil.copy2(source, interview_dir / source.name)
    source_screenshot_dir = PROJECT_ROOT / "docs" / "screenshots"
    if source_screenshot_dir.exists():
        for source in source_screenshot_dir.glob("*.*"):
            if source.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
                shutil.copy2(source, screenshot_dir / source.name)
    combined = "\n\n---\n\n".join(markdown for _, markdown in contents)
    (destination / "项目完整说明.txt").write_text(plain_text(combined), encoding="utf-8")
    build_docx(contents, word_dir / "审脉AuditPilot_项目完整说明.docx")
    build_pdf(contents, pdf_dir / "审脉AuditPilot_项目完整说明.pdf")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--destination", required=True, type=Path)
    args = parser.parse_args()
    build(args.destination.resolve())
    print(args.destination.resolve())


if __name__ == "__main__":
    main()
